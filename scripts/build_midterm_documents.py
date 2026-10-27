from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "docx"
SCRIPT_PATH = OUTPUT_DIR / "MLA-2026-midterm-presentation-script.docx"
QNA_PATH = OUTPUT_DIR / "MLA-2026-midterm-qna.docx"

NAVY = "183153"
PALE_BLUE = "EAF1F6"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = RGBColor(85, 85, 85)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_GRAY, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=11, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color


def clear_style_paragraph_decoration(style):
    p_pr = style.element.get_or_add_pPr()
    for tag_name in ("w:pBdr", "w:shd"):
        child = p_pr.find(qn(tag_name))
        if child is not None:
            p_pr.remove(child)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(6)
    clear_style_paragraph_decoration(title)

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = "Calibri"
    heading1.font.size = Pt(16)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(0, 0, 0)
    heading1.paragraph_format.space_before = Pt(10)
    heading1.paragraph_format.space_after = Pt(5)
    heading1.paragraph_format.keep_with_next = True

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = "Calibri"
    heading2.font.size = Pt(13)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(0, 0, 0)
    heading2.paragraph_format.space_before = Pt(8)
    heading2.paragraph_format.space_after = Pt(4)
    heading2.paragraph_format.keep_with_next = True


def add_title(doc, title_text, subtitle_text):
    title = doc.add_paragraph(title_text, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run(subtitle_text)
    set_run_font(run, size=11, italic=True, color=TEXT_GRAY)


def add_label_paragraph(doc, label, text, after=6):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run)
    return paragraph


SCRIPT_SECTIONS = [
    {
        "slide": "Slide 1  Project title",
        "time": "0:00 to 0:25",
        "script": (
            "Good morning. Our project is called CEFR Based Sentence Difficulty Classification for "
            "Teacher Support in Reading Material Review. We propose a support tool that predicts the "
            "CEFR difficulty of individual English sentences. The goal is to help a teacher review reading "
            "materials more consistently while keeping every final decision with the teacher."
        ),
        "cue": "Start slowly. Make eye contact before moving to the problem.",
    },
    {
        "slide": "Slide 2  Problem and user",
        "time": "0:25 to 1:15",
        "script": (
            "At HANU, teachers often select or adapt reading texts for a specific class. A passage may look "
            "suitable overall but still contain a few sentences that are much harder than the target level. "
            "Checking every sentence manually takes time, and reviewers may judge the same sentence differently. "
            "Our tool accepts a teacher selected reading text, separates it into sentences, and returns a predicted "
            "CEFR level with six class probabilities for each sentence. The teacher can then keep, simplify, replace, "
            "or examine a sentence more closely. The tool supports material review. It does not grade learners or "
            "choose materials automatically."
        ),
        "cue": "Emphasize the final sentence because it establishes human control.",
    },
    {
        "slide": "Slide 3  Data",
        "time": "1:15 to 2:20",
        "script": (
            "We plan to use UniversalCEFR cefr sp en, matched to the official CEFR-SP files. It contains 10,004 "
            "independent English sentences from 5 to 34 words, with a median of 14 and a mean of about 14.8. The fields "
            "include text, CEFR level, source, category, format, and license. B1 and B2 contain most of the data, while A1 "
            "has 124 sentences and C2 has 230. The sources are SCoRE and Wiki-Auto. Hugging Face marks the data as CC "
            "BY-NC-SA 4.0. The original repository lists SCoRE under that license and Wiki-Auto under CC BY-SA 3.0. One "
            "real example is Is that your bike, labeled A1. CEFR-SP provides two expert labels. Their exact agreement is "
            "about 41.4 percent, although every disagreement is within one adjacent CEFR level. These labels therefore "
            "require cautious interpretation."
        ),
        "cue": "Point briefly to the class distribution. Avoid reading every bar value.",
    },
    {
        "slide": "Slide 4  ML task alignment",
        "time": "2:20 to 3:15",
        "script": (
            "We formulate the project as six class sentence classification. The input is one independent English "
            "sentence. The output is a probability for A1, A2, B1, B2, C1, and C2, followed by the predicted level. "
            "For example, the input Is that your bike has the target label A1. CEFR levels have an order, so an A1 "
            "prediction of A2 is less serious than a prediction of C2. We still train a six class classifier, but our "
            "evaluation also measures the distance between the predicted and true levels. Showing probabilities is "
            "important because the labels are not perfectly objective and the teacher needs to see uncertainty."
        ),
        "cue": "Pause on the A1 to A2 and A1 to C2 comparison.",
    },
    {
        "slide": "Slide 5  Model plan",
        "time": "3:15 to 4:25",
        "script": (
            "Our baseline always predicts B2, the most frequent class in the training data. This is the simple method "
            "that both trained models should beat. Model one uses word unigram and bigram TF-IDF features with "
            "multinomial Logistic Regression and balanced class weights. We expect vocabulary and short phrases to "
            "provide strong CEFR signals. We will inspect its top coefficients and compare performance by source because "
            "TF-IDF could learn topic or dataset source instead of difficulty. Model two is a BiLSTM with 100 dimensional "
            "GloVe embeddings, dropout, class weighted cross entropy, and early stopping. It tests whether word order and "
            "sequential structure add useful information. The sentences are short and the dataset has about 10,000 "
            "examples, so the BiLSTM may produce only a small gain or may overfit. That limitation is part of the comparison."
        ),
        "cue": "Do not imply that the BiLSTM will win. Present it as a testable hypothesis.",
    },
    {
        "slide": "Slide 6  Evaluation plan",
        "time": "4:25 to 5:50",
        "script": (
            "We will use the official CEFR-SP split: 7,178 sentences for training, 1,366 for validation, and 1,460 for "
            "testing. The split was designed to keep highly similar sentences from appearing in both training and "
            "validation or test data. We will remove one exact duplicate inside the training set, leaving 7,177 unique "
            "training sentences. We will fit the tokenizer, vocabulary, and TF-IDF representation on training data only. "
            "The validation set will guide hyperparameters and early stopping, and the test set will be used once after "
            "the model choices are fixed. Our primary metric is macro-F1 because it gives equal importance to every class. "
            "We will also report quadratic weighted kappa because CEFR levels are ordered and distant errors matter more. "
            "The report will include accuracy, a confusion matrix, and per-class precision, recall, F1, and support. We will "
            "estimate 95 percent confidence intervals with 1,000 stratified bootstrap samples. A1 results need cautious "
            "interpretation because the test set contains only 37 A1 sentences."
        ),
        "cue": "Spend most of this slide on split discipline and the reasons for the two main metrics.",
    },
    {
        "slide": "Slide 7  Risks and responsible use",
        "time": "5:50 to 6:40",
        "script": (
            "The main risks are label ambiguity, class imbalance, topic confounding, and a mismatch between the public "
            "data and HANU reading materials. We address imbalance with class weighting and macro-F1. We check topic "
            "confounding through Logistic Regression coefficients and results by source. Before the final project, we will "
            "also review about 30 representative sentences from HANU materials as a qualitative local check. These sentences "
            "will not be used for training. The model evaluates sentences, not complete passages, and the same sentence may "
            "feel different to learners with different background knowledge. For these reasons, predictions remain suggestions "
            "with probabilities. The teacher reviews the sentence in context and makes the final decision. Thank you. We are "
            "ready for your questions."
        ),
        "cue": "Finish before 6:40 to preserve a timing buffer inside the seven minute limit.",
    },
]


QNA_GROUPS = [
    (
        "Project and data",
        [
            (
                "1  Why classify individual sentences instead of complete passages",
                "The dataset labels individual sentences, so sentence classification is the task we can support directly. It also gives teachers an actionable result because they can see which sentence may need attention. Passage difficulty depends on discourse, cohesion, and context across sentences, which this model does not capture. We therefore avoid assigning one CEFR label to an entire passage.",
            ),
            (
                "2  Why use CEFR as the target",
                "CEFR gives teachers a familiar ordered scale from A1 to C2 and fits the practical decision of reviewing material for a target level. It is still an expert judgment rather than a perfect physical measurement. That is why the tool shows probabilities and keeps the teacher in control.",
            ),
            (
                "3  Is 41.4 percent exact annotator agreement too low",
                "It shows that sentence difficulty is genuinely ambiguous. However, every public disagreement is within one adjacent CEFR level, so the experts usually disagree about a boundary rather than placing a sentence at opposite ends of the scale. We will report this limitation, use an ordinal metric, and present predictions as suggestions rather than facts.",
            ),
            (
                "4  Why does the Hugging Face dataset use the higher expert label",
                "In our row-by-row match, the Hugging Face label corresponds to the higher of the two CEFR-SP expert labels. We are reproducing the published data mapping rather than choosing a label after seeing model performance. We will document this choice and retain the two-label information for error analysis where possible.",
            ),
            (
                "5  Which licences apply to the dataset",
                "Hugging Face marks the dataset as CC BY-NC-SA 4.0. The original repository lists SCoRE under CC BY-NC-SA 4.0 and Wiki-Auto under CC BY-SA 3.0. We therefore limit the project to non-commercial educational research and retain the source, attribution, and licence metadata for every record.",
            ),
        ],
    ),
    (
        "Models and baseline",
        [
            (
                "6  Why is the baseline always B2",
                "B2 is the most frequent class in the training set, so always predicting B2 is the standard majority-class baseline. It shows what a model can achieve without learning any relationship between language and difficulty. Accuracy against this baseline is useful, but macro-F1 is necessary because the baseline performs poorly on rare classes.",
            ),
            (
                "7  Why is Logistic Regression not the baseline",
                "Logistic Regression with TF-IDF is already a trained model with thousands of learned feature weights. A baseline should be simpler and require almost no learning. We therefore separate the majority-class predictor from the two model families we want to compare.",
            ),
            (
                "8  Why should TF-IDF work for CEFR difficulty",
                "CEFR levels often differ in vocabulary, phrase choice, and the frequency of more advanced words. Unigrams and bigrams can capture many of those signals, and Logistic Regression is interpretable. Its weakness is that it may learn subject matter or source patterns, so we will inspect coefficients and compare performance across SCoRE and Wiki-Auto.",
            ),
            (
                "9  Why use a BiLSTM when the sentences are short",
                "The BiLSTM tests whether word order and sequential structure add information beyond unordered TF-IDF features. We do not assume that it will outperform Logistic Regression. Short sentences and about 10,000 examples may limit its advantage, which makes the comparison a useful empirical question.",
            ),
            (
                "10  Why not use a transformer model",
                "The current plan focuses on two model families covered by the course and keeps the comparison feasible for the available data and time. A transformer could become an extension after we establish reliable baselines and evaluation. Adding it now would increase compute and tuning requirements without solving the label and domain limitations.",
            ),
        ],
    ),
    (
        "Evaluation",
        [
            (
                "11  How will you prevent data leakage",
                "We will keep the official CEFR-SP train, validation, and test sets. The authors designed this split so that highly similar sentences do not cross from training into validation or test. We will fit the tokenizer, vocabulary, and TF-IDF features only on training data, and we will use the test set once after model selection.",
            ),
            (
                "12  Why is accuracy insufficient",
                "The classes are strongly imbalanced. A system can gain a reasonable accuracy by favoring B1 and B2 while performing badly on A1 and C2. Macro-F1 gives each CEFR class equal influence, so it reveals whether the model works across the full scale.",
            ),
            (
                "13  Why use quadratic weighted kappa",
                "CEFR labels are ordered. Predicting A2 for an A1 sentence is a smaller error than predicting C2. Quadratic weighted kappa gives larger penalties to errors that are farther apart, so it adds information that ordinary accuracy and macro-F1 do not provide.",
            ),
            (
                "14  Can you trust the A1 result with only 37 test examples",
                "Only with caution. We will report the support beside the class metric and use stratified bootstrap samples to estimate uncertainty. We will avoid strong claims based on A1 alone. A larger labeled local sample would be needed before using the model for high-stakes decisions about beginner materials.",
            ),
        ],
    ),
    (
        "Risks and use",
        [
            (
                "15  How will you know whether the model transfers to HANU materials",
                "The public data cannot prove local validity. Before the final project, we plan a qualitative check with about 30 representative sentences from HANU reading materials. Teachers can inspect whether the predicted levels and uncertainty are plausible. These sentences are for local review only and will not enter training.",
            ),
            (
                "16  What happens if the BiLSTM performs worse than Logistic Regression",
                "That is still a valid result. It would suggest that lexical features are sufficient for this dataset or that the sequence model needs more data. We would report the comparison honestly, analyze the errors, and prefer the simpler model if it performs as well or better.",
            ),
            (
                "17  Could the tool replace a teacher's judgment",
                "No. The model sees one sentence and cannot know the lesson objective, the learners' background knowledge, or how the sentence functions inside the passage. It provides a predicted level and probabilities to focus attention. The teacher reviews the sentence in context and makes the final decision.",
            ),
        ],
    ),
]


def build_script_document():
    doc = Document()
    configure_document(doc)
    add_title(
        doc,
        "MLA 2026 Midterm Presentation Script",
        "One presenter  Seven slides  Target delivery time 6 minutes 40 seconds",
    )

    intro = doc.add_paragraph(
        "This script follows the Week 6 brief and covers the six required project areas. "
        "The timing leaves about 20 seconds inside the strict seven minute limit."
    )
    intro.paragraph_format.space_after = Pt(10)

    doc.add_heading("Timing plan", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(0.75), Inches(4.45), Inches(1.45)]
    headers = ["Slide", "Topic", "Target"]
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(headers[index])
        set_run_font(run, size=10, bold=True, color=RGBColor(255, 255, 255))
    set_repeat_table_header(table.rows[0])

    for number, section in enumerate(SCRIPT_SECTIONS, 1):
        row = table.add_row()
        values = [str(number), section["slide"].split("  ", 1)[1], section["time"]]
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if number % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index != 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(values[index])
            set_run_font(run, size=10)

    doc.add_paragraph()
    doc.add_heading("Full script", level=1)

    for section in SCRIPT_SECTIONS:
        doc.add_heading(section["slide"], level=2)
        time_p = doc.add_paragraph()
        time_p.paragraph_format.space_after = Pt(4)
        time_run = time_p.add_run(f"Target time  {section['time']}")
        set_run_font(time_run, size=10, italic=True, color=TEXT_GRAY)

        script_p = doc.add_paragraph(section["script"])
        script_p.paragraph_format.space_after = Pt(4)
        script_p.paragraph_format.keep_together = False

        cue_p = doc.add_paragraph()
        cue_p.paragraph_format.space_after = Pt(8)
        cue_label = cue_p.add_run("Delivery cue: ")
        set_run_font(cue_label, size=10, bold=True, italic=True, color=TEXT_GRAY)
        cue_text = cue_p.add_run(section["cue"])
        set_run_font(cue_text, size=10, italic=True, color=TEXT_GRAY)

    doc.add_heading("Source notes", level=1)
    sources = [
        "UniversalCEFR cefr_sp_en  https://huggingface.co/datasets/UniversalCEFR/cefr_sp_en",
        "CEFR-SP repository  https://github.com/yukiar/CEFR-SP/tree/main/CEFR-SP",
        "Arase et al. 2022  https://aclanthology.org/2022.emnlp-main.416/",
    ]
    for source in sources:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(source)
        set_run_font(run, size=10)

    doc.core_properties.title = "MLA 2026 Midterm Presentation Script"
    doc.core_properties.subject = "Seven slide script for the Week 6 final project idea presentation"
    doc.save(SCRIPT_PATH)


def build_qna_document():
    doc = Document()
    configure_document(doc)
    add_title(
        doc,
        "MLA 2026 Midterm Questions and Answers",
        "Moderate difficulty practice for the three minute panel question period",
    )

    intro = doc.add_paragraph(
        "Use these answers as speaking guides. Start with the direct answer, give one reason or number, "
        "and stop after about 20 to 40 seconds unless the panel asks a follow-up question."
    )
    intro.paragraph_format.space_after = Pt(10)

    for group_title, questions in QNA_GROUPS:
        doc.add_heading(group_title, level=1)
        for question, answer in questions:
            doc.add_heading(question, level=2)
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            label = paragraph.add_run("Suggested answer: ")
            set_run_font(label, bold=True)
            body = paragraph.add_run(answer)
            set_run_font(body)

    doc.add_heading("Source notes", level=1)
    sources = [
        "UniversalCEFR cefr_sp_en  https://huggingface.co/datasets/UniversalCEFR/cefr_sp_en",
        "CEFR-SP repository  https://github.com/yukiar/CEFR-SP/tree/main/CEFR-SP",
        "Arase et al. 2022  https://aclanthology.org/2022.emnlp-main.416/",
    ]
    for source in sources:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(source)
        set_run_font(run, size=10)

    doc.core_properties.title = "MLA 2026 Midterm Questions and Answers"
    doc.core_properties.subject = "Moderate difficulty panel questions and suggested answers"
    doc.save(QNA_PATH)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_script_document()
    build_qna_document()
    print(SCRIPT_PATH)
    print(QNA_PATH)


if __name__ == "__main__":
    main()
