"""Create plain student-style documents from the verified baseline outputs."""

import csv
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
EXPERIMENT = ROOT / "output" / "experiments" / "sentence-baselines"
OUTPUT = ROOT / "output" / "docx"
REFERENCE = ROOT / "docs" / "reference" / "MLA-2026.docx"
LABELS = ["A1", "A2", "B1", "B2", "C1", "C2"]
NAMES = {"majority": "Always B2", "awl": "AWL density only", "tfidf_lr": "TF-IDF + LR"}


def load_json(name):
    return json.loads((EXPERIMENT / name).read_text(encoding="utf-8"))


def style_setup(style, size, bold=False, before=0, after=4):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    props = style.element.get_or_add_rPr()
    color = props.find(qn("w:color"))
    if color is not None:
        for key in ["w:themeColor", "w:themeTint", "w:themeShade"]:
            color.attrib.pop(qn(key), None)
    pformat = style.paragraph_format
    pformat.space_before = Pt(before)
    pformat.space_after = Pt(after)
    pformat.line_spacing = 1.0
    pformat.widow_control = True
    for tag in ["w:pBdr", "w:shd"]:
        element = style.element.get_or_add_pPr().find(qn(tag))
        if element is not None:
            element.getparent().remove(element)


def plain_doc(body_size=11):
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == "f797e0fb90f7a9bae1c3a83032d845a8fa74abe1c69dd0e76d931ab05fde1171"
    doc = Document(REFERENCE)
    body = doc.element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    for section in doc.sections:
        section.page_width = Inches(8.2677)
        section.page_height = Inches(11.6929)
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
    style_setup(doc.styles["Normal"], body_size, after=5)
    style_setup(doc.styles["Title"], 16, True, after=5)
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_setup(doc.styles["Subtitle"], 10.5, after=7)
    doc.styles["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_setup(doc.styles["Heading 1"], 12, True, before=8, after=4)
    style_setup(doc.styles["Heading 2"], 11.5, True, before=7, after=3)
    doc.core_properties.author = "Group 10 CLC02"
    doc.core_properties.subject = "CEFR sentence classification project"
    return doc


def paragraph(doc, text, after=None):
    p = doc.add_paragraph(text)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def heading(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.keep_with_next = True
    return p


def hyperlink(paragraph, label, url):
    relation = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def plain_table(doc, headers, rows, widths, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    total = sum(widths)
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "BFBFBF")
        borders.append(element)
    margins = OxmlElement("w:tblCellMar")
    for side, value in [("top", "55"), ("bottom", "55"), ("left", "75"), ("right", "75")]:
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tblpr.append(margins)
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = str(text)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, row in enumerate(table.rows):
        keep = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(keep)
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                border = OxmlElement(f"w:{edge}")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:color"), "BFBFBF")
                cell_borders.append(border)
            cell._tc.get_or_add_tcPr().append(cell_borders)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "FFFFFF")
            shading.set(qn("w:val"), "clear")
            cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(font_size)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
    return table


def page(doc, title):
    doc.add_page_break()
    heading(doc, title)


def proposal(results):
    doc = plain_doc(11)
    doc.add_paragraph("CEFR Sentence Classification for Reading Material Adaptation", style="Title")
    doc.add_paragraph("Group 10 CLC02   |   Project proposal   |   4 October 2026", style="Subtitle")
    heading(doc, "1 Problem and User")
    paragraph(doc, "We propose a tool for HANU English teachers who want to simplify reading materials for a class with a target CEFR level. It highlights sentences predicted above that level so the teacher can review, simplify or replace them.")
    heading(doc, "2 Data")
    paragraph(doc, "We use 10,004 public CEFR-SP sentence records from SCoRE and Wiki-Auto, available through UniversalCEFR [1-3]. Our target is the higher expert label, matching the Hugging Face version. We retain official splits, both annotations and source licenses: SCoRE CC BY-NC-SA 4.0; Wiki-Auto CC BY-SA 3.0. Use is non-commercial educational research.")
    heading(doc, "3 ML Task Alignment")
    paragraph(doc, "Input: one English sentence. Output: a label from A1, A2, B1, B2, C1 and C2, with six model probabilities. In use, the teacher supplies the class level and reviews highlighted sentences in context. The tool does not assign a CEFR level to the whole passage or rewrite it automatically.")
    heading(doc, "4 Model Plan")
    paragraph(doc, "Baselines: always predict B2; classify from Academic Word List (AWL) density [4] using a small decision tree; and word unigram-bigram TF-IDF with balanced multinomial Logistic Regression. The next model is a BiLSTM with GloVe 100-dimensional embeddings, dropout, weighted loss and validation early stopping. It tests whether word order adds information beyond lexical features.")
    heading(doc, "5 Evaluation Plan")
    paragraph(doc, "Official splits: train 7,177 after deduplication; validation 1,366; test 1,460. Vocabulary and TF-IDF are fitted on train; model settings are selected on validation before test evaluation. QWK measures ordered agreement; MAE measures average CEFR-rank error. Per-class support and F1 expose minority-class failures. Current test results are below; details accompany this proposal.")
    plain_table(doc, ["Baseline", "QWK higher is better", "MAE lower is better"], [
        [NAMES[name], f"{score['qwk']:.3f}", f"{score['mae']:.3f}"]
        for name, score in results["models"].items()
    ], [2.87, 1.7, 1.7], font_size=10.5)
    heading(doc, "6 Risks and Responsible Use")
    paragraph(doc, "LR has zero A1 and C2 F1; test support is 37 and 72. The selected AWL baseline is unreliable. Expert disagreement, source cues and differences from HANU materials limit interpretation. Probabilities are not calibrated confidence. Teachers keep final control. A qualitative review of about 30 HANU sentences is planned; classroom usefulness is not yet established.")
    p = paragraph(doc, "Sources: ", after=0)
    p.paragraph_format.space_before = Pt(6)
    for i, (label, url) in enumerate([
        ("[1] UniversalCEFR", "https://huggingface.co/datasets/UniversalCEFR/cefr_sp_en"),
        ("[2] CEFR-SP files", "https://github.com/yukiar/CEFR-SP/tree/main/CEFR-SP"),
        ("[3] Arase et al 2022", "https://aclanthology.org/2022.emnlp-main.416/"),
        ("[4] Coxhead AWL", "https://www.wgtn.ac.nz/lals/resources/academicwordlist"),
    ]):
        if i:
            p.add_run("; ")
        hyperlink(p, label, url)
    for run in p.runs:
        run.font.size = Pt(9.5)
    doc.core_properties.title = "CEFR Sentence Classification for Reading Material Adaptation"
    target = OUTPUT / "Group10_CLC02_one_page_proposal.docx"
    doc.save(target)
    return target


def baseline_report(results, selected, audit, eda):
    doc = plain_doc(11)
    doc.add_paragraph("CEFR Sentence Baseline Results", style="Title")
    doc.add_paragraph("Group 10 CLC02   |   Experiment run on 4 October 2026", style="Subtitle")
    paragraph(doc, "TF-IDF with Logistic Regression performs best among the three tested baselines on QWK and MAE. It still fails to identify A1 and C2 correctly. These results support further sentence-level experiments, not a claim that the tool is ready for classroom use.")

    heading(doc, "Evaluation setup")
    paragraph(doc, "The public 10,004-record dataset is matched to the six official SCoRE and Wiki-Auto files. Targets are the higher of the two expert labels, as in UniversalCEFR. The official split has 7,178 train, 1,366 validation and 1,460 test records. Removing one identical B1 training copy leaves 7,177 train records. Exact sentence overlap across splits is zero.")
    paragraph(doc, "Models are fitted on train and selected by highest validation QWK, then lower MAE and higher macro-F1 for ties. Selection is frozen before the test phase. The following table uses the same 1,460 test sentences for all models. BiLSTM has not been trained.")
    heading(doc, "Test results")
    rows = []
    for name, score in results["models"].items():
        if name == "majority":
            qwk, mae = f"{score['qwk']:.3f}", f"{score['mae']:.3f}"
        else:
            qwk = f"{score['qwk']:.3f} ({score['ci95']['qwk'][0]:.3f}-{score['ci95']['qwk'][1]:.3f})"
            mae = f"{score['mae']:.3f} ({score['ci95']['mae'][0]:.3f}-{score['ci95']['mae'][1]:.3f})"
        rows.append([NAMES[name], qwk, mae])
    plain_table(doc, ["Baseline", "QWK and 95% CI", "MAE and 95% CI"], rows, [2.47, 1.9, 1.9])
    paragraph(doc, "QWK is chance-corrected agreement with larger penalties for distant errors; higher is better. MAE is the average absolute rank difference using A1=0 through C2=5; lower is better. Rank steps are an evaluation convention, not equal measured increments of language ability. AWL has positive QWK but substantially worse MAE than always predicting B2.")

    heading(doc, "LR performance by class")
    plain_table(doc, ["Class", "Test n", "Predicted n", "Precision", "Recall", "F1"], [
        [row["label"], row["support"], row["predicted_count"], f"{row['precision']:.3f}", f"{row['recall']:.3f}", f"{row['f1']:.3f}"]
        for row in results["models"]["tfidf_lr"]["per_class"]
    ], [0.63, 0.71, 1.0, 1.31, 1.31, 1.31])
    paragraph(doc, "A1 has no correct predictions; C2 is never predicted. Balanced class weights do not resolve these failures. As a supplementary imbalance diagnostic, macro-F1 is 0.266 for LR, 0.078 for AWL and 0.080 for always B2. Accuracy is retained in the experiment data but is not a primary criterion.")

    page(doc, "Data analysis on the training set")
    paragraph(doc, "The table below uses only the deduplicated training set for feature descriptions. Alphabetic word tokens exclude numbers and punctuation, keep attached apostrophes and hyphens, and are lowercased. AWL density counts repeated occurrences of listed AWL forms. A rare word here means document frequency at most five in train; this is a corpus-relative diagnostic, not a measure of learner familiarity.")
    training = [row for row in eda if row["split"] == "train"]
    plain_table(doc, ["Class", "Train n", "Mean words", "Mean AWL %", "Zero AWL %", "Rare tokens %"], [
        [row["cefr_level"], row["support"], f"{float(row['word_count_mean']):.1f}", f"{100*float(row['awl_density_mean']):.1f}",
         f"{100*float(row['zero_awl_fraction']):.1f}", f"{100*float(row['train_rare_word_ratio_mean']):.1f}"]
        for row in training
    ], [0.6, 0.8, 1.08, 1.2, 1.19, 1.4])
    paragraph(doc, "A-level mean lengths are shorter than those of higher levels, but the C1 and C2 means are close. Mean AWL density rises to B2, then decreases in C2. Almost half of C2 training sentences contain no AWL forms. This contradicts a simple rule that more academic vocabulary always means a higher CEFR level.")

    heading(doc, "Why the rare classes are difficult")
    paragraph(doc, "A1 has only 47 training examples and C2 only 91. The A1 source composition also changes: 41 of 47 training examples are from Wiki-Auto, while 33 of 37 test examples are from SCoRE. The largest A1 LR coefficients include 'died on' and 'died', which may reflect short biography-like sentences rather than broadly transferable A1 language. C2 has more corpus-rare vocabulary and few training examples. These are plausible explanations, not proven causes of the errors.")
    heading(doc, "LR results by source")
    plain_table(doc, ["Test source", "Test n", "QWK", "MAE"], [
        [source, score["n"], f"{score['qwk']:.3f}", f"{score['mae']:.3f}"]
        for source, score in results["models"]["tfidf_lr"]["per_source"].items()
    ], [2.17, 1.0, 1.5, 1.6])
    paragraph(doc, "The source class mixes are different: SCoRE has no C2 test examples and only five C1 examples; Wiki-Auto contains all 72 C2 test examples. Pooled QWK is higher than either source-specific QWK, so between-source differences may contribute to the apparent performance. The values are not directly comparable without considering their label distributions. Coefficients such as 'potassium' and 'enzyme' warrant checking for topic cues.")

    heading(doc, "Length and label conventions")
    paragraph(doc, "Across all 10,004 records, whitespace splitting gives 5-34 tokens, median 14 and mean 14.8, including punctuation and numbers. The alphabetic-word convention used here gives 2-31 tokens, median 12 and mean 13.0. These are different counting units. Exact expert agreement is 41.43% in this filtered public subset; the other pairs differ by one level. That is not the raw agreement rate before corpus filtering.")

    page(doc, "Baseline definitions and selection")
    heading(doc, "Majority predictor")
    paragraph(doc, "The most frequent training target is B2, with 2,472 of 7,177 examples. This predictor always outputs B2 and provides a minimal comparison, separate from trained classifiers.")
    heading(doc, "AWL density baseline")
    paragraph(doc, "The source is Coxhead's Academic Word List from Victoria University of Wellington [4]. Its 570 families and ten sublists concern academic word frequency, not CEFR levels. The full family-list PDF yields 3,111 distinct listed forms; all 570 headwords are verified against it. Both British and American variants are included.")
    paragraph(doc, "Feature: the count of word tokens matching any listed AWL form divided by the sentence's alphabetic word-token count. Matching is case-insensitive and includes repeated occurrences. No stemming or guessed family members are added. A six-class decision tree learns density thresholds from train; it does not import a CEFR mapping from AWL.")
    paragraph(doc, "Eighteen settings are tested on validation: maximum leaves 2, 3 or 6; minimum leaf size 50, 100 or 200; and no class weighting or balanced weighting. Highest validation QWK selects six maximum leaves, minimum leaf size 100 and balanced weights. The fitted model uses four predicted labels and never predicts A2 or B1.")
    plain_table(doc, ["Learned density interval", "Predicted label"], [
        ["0 to 0.017241", "A1"],
        ["Above 0.017241 to 0.048810", "C2"],
        ["Above 0.048810 to 0.270979", "C1"],
        ["Above 0.270979", "B2"],
    ], [4.47, 1.8])
    paragraph(doc, "Thresholds are rounded for display; adjacent leaves with the same predicted class are combined here. The saved fitted tree is the authoritative rule. Its non-monotonic mapping and large A1 prediction count show why the feature cannot be presented as a CEFR scale. The selected tree predicts A1 for 836 test sentences although only 37 test targets are A1.")
    paragraph(doc, "Selection has a tradeoff: an unweighted two-leaf tree has validation QWK 0.224 and MAE 0.769; the selected weighted tree has QWK 0.228 and MAE 1.863. Only the selected tree is tested. This documents a limitation of choosing solely by highest QWK, not a reason to replace it retrospectively using test results. AWL can remain a diagnostic feature; this implementation is not a reliable standalone classifier.")

    heading(doc, "TF IDF with Logistic Regression")
    paragraph(doc, "Word unigrams and bigrams use lowercasing, sublinear term frequency, L2 normalization, at most 30,000 features and no stop-word removal. The vectorizer's default pattern keeps alphanumeric tokens of at least two characters. Vocabulary and inverse document frequencies are fitted on train only. Multinomial LR uses L2 regularization, balanced train-derived weights, the lbfgs solver and a 2,000-iteration limit [5].")
    paragraph(doc, "Six validation settings are tried: C in {0.1, 1, 10} and minimum document frequency in {1, 2}. Selected settings are C=10 and min_df=1, with 30,000 features. Training converges in 160 iterations. Validation QWK is 0.623 and MAE 0.693. The model is evaluated on test without refitting on train plus validation.")

    page(doc, "Interpretation and next experiments")
    heading(doc, "What the experiment establishes")
    paragraph(doc, "On this public sentence benchmark, TF-IDF with LR improves on always B2 in both primary metrics. AWL density alone does not provide an adequate six-class classifier under the selected rule. Macro-F1 and per-class counts reveal failures that the overall ordinal scores do not resolve. No result assigns a passage-level CEFR label or measures the quality of a rewritten text.")
    heading(doc, "Uncertainty and limitations")
    paragraph(doc, "The 95% intervals use 1,000 bootstrap samples, stratified by true class with paired row indices across models. Fitted models and test class proportions are held fixed. For the constant B2 predictor, the conditional intervals collapse to its point values because every resample keeps the same class counts. These intervals do not cover training-seed variability or changes to the classroom population.")
    paragraph(doc, "The maximum-expert target convention is stricter than accepting either expert label. The original CEFR-SP study uses both labels and also includes the non-public Newsela portion [3]. These scores therefore are not a direct reproduction or comparison with its published model scores. A C1 prediction for a maximum-label C2 sentence can sometimes match the other annotator, which should be considered during qualitative error review.")
    paragraph(doc, "The official split was designed to reduce train-to-evaluation sentence similarity [3]. The exact-overlap audit does not prove all document-level leakage is absent, because document identifiers are not available in this public representation. SCoRE and Wiki-Auto also differ from HANU materials. Model probabilities, particularly after class weighting, must not be described as calibrated confidence or as the probability that every learner will find a sentence difficult.")
    heading(doc, "Next work")
    for text in [
        "Train the planned BiLSTM with GloVe 100-dimensional embeddings, dropout, class-weighted loss and early stopping. Select settings on validation and compare using the same target definition, QWK, MAE and per-class support. A gain over LR remains a hypothesis.",
        "Review coefficient and source diagnostics. Reserve the test set for final comparisons, rather than repeatedly choosing new rules against its errors. Changes suggested by this exploratory analysis should be developed using train and validation.",
        "Review about 30 representative HANU sentences with a teacher. For a teacher-specified class level, examine both highlighted and unhighlighted sentences, especially missed difficult content. This qualitative check is planned, not completed, and cannot establish broad learner validity.",
    ]:
        paragraph(doc, text)
    heading(doc, "References and reproducibility")
    entries = [
        ("[1] UniversalCEFR cefr_sp_en", "https://huggingface.co/datasets/UniversalCEFR/cefr_sp_en"),
        ("[2] CEFR-SP official data files and licenses", "https://github.com/yukiar/CEFR-SP/tree/main/CEFR-SP"),
        ("[3] Arase, Uchida and Kajiwara (2022). CEFR-Based Sentence Difficulty Annotation and Assessment. EMNLP, pp. 6206-6219.", "https://aclanthology.org/2022.emnlp-main.416/"),
        ("[4] Coxhead Academic Word List and full family forms", "https://www.wgtn.ac.nz/lals/resources/academicwordlist/publications/awlsublists1.pdf"),
        ("[5] scikit-learn 1.7 LogisticRegression and evaluation documentation", "https://scikit-learn.org/1.7/modules/generated/sklearn.linear_model.LogisticRegression.html"),
    ]
    for label, url in entries:
        p = paragraph(doc, "", after=4)
        hyperlink(p, label, url)
    paragraph(doc, "The project retains raw-data hashes, the validation grid, frozen models, per-sentence predictions, all confusion matrices and per-class reports. The reproducibility notes and experiment files are in output/experiments/sentence-baselines. Python 3.12.14 and scikit-learn 1.7.2 are used. Source licenses are retained; this is non-commercial educational research.")
    doc.core_properties.title = "CEFR Sentence Baseline Results"
    target = OUTPUT / "Group10_CLC02_sentence_baseline_results.docx"
    doc.save(target)
    return target


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    results = load_json("test_results.json")
    selected = load_json("frozen_selection.json")
    audit = load_json("data_audit.json")
    assert selected["test_evaluated"]
    assert audit["used_split_counts"] == {"train": 7177, "dev": 1366, "test": 1460}
    with (EXPERIMENT / "eda_by_level.csv").open(encoding="utf-8-sig", newline="") as stream:
        eda = list(csv.DictReader(stream))
    paths = [proposal(results), baseline_report(results, selected, audit, eda)]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
