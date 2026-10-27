from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx" / "MLA-2026-revised-outline.docx"

BLACK = "000000"
LIGHT_GRAY = "E7E6E6"
BORDER_GRAY = "D9D9D9"
DARK = BLACK
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=110, bottom=60, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="000000", size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_table(doc, headers, rows, widths=None, font_size=10.5, keep_whole=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(headers) > 3 and i > 0 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(BLACK)
        run.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(headers) > 3 and i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor.from_string(BLACK)
    if keep_whole:
        for table_row in table.rows[:-1]:
            for cell in table_row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    return table


def add_labeled_paragraph(doc, label, text, *, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = keep
    r = p.add_run(label)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLACK)
    p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.45 + 0.45 * level)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, name="Calibri", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc):
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Cm(21.01)
    section.page_height = Cm(29.69)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(5)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name in ("Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    doc.styles["Heading 1"].font.size = Pt(20)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    doc.styles["Heading 2"].font.size = Pt(16)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)

    for style_name in ("List Bullet", "List Bullet 2"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(12)

    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_after = Pt(0)

    footer = section.footer
    footer.paragraphs[0].text = ""


def build_document():
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("CEFR Based Sentence Difficulty Classification\nfor Teacher Support in Reading Material Review")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("Phân loại độ khó câu theo CEFR để hỗ trợ giáo viên rà soát tài liệu đọc")
    set_run_font(run, size=20, color=BLACK, bold=True)

    add_labeled_paragraph(
        doc,
        "Scope. ",
        "One English sentence in; six CEFR probabilities and one predicted level out; the teacher decides what to do.",
    )

    doc.add_heading("1. Problem and User", level=1)
    add_labeled_paragraph(
        doc,
        "Problem. ",
        "Teachers reviewing English reading materials must inspect individual sentences to identify language that may exceed a target CEFR level. This sentence-by-sentence review is time-consuming, especially when a text contains only a few potentially unsuitable sentences.",
    )
    add_labeled_paragraph(
        doc,
        "Primary user. ",
        "English-language teachers at HANU who select or adapt reading materials for a known learner group.",
    )
    add_labeled_paragraph(
        doc,
        "Decision supported. ",
        "For each sentence, the tool predicts A1, A2, B1, B2, C1, or C2 and displays the probability distribution. The teacher can then decide whether to keep, simplify, replace, or inspect the sentence more closely.",
    )
    add_labeled_paragraph(
        doc,
        "Operational workflow. ",
        "A teacher pastes or uploads a selected reading text; the interface segments it into sentences; the classifier analyzes each English sentence independently; and the interface highlights sentences whose predicted level may not match the teacher's target. This is how a sentence-level training dataset can support review of a teacher-selected reading text.",
    )
    add_labeled_paragraph(
        doc,
        "Scope boundary and human control. ",
        "The project does not assign a single automated CEFR rating to an entire reading text. Predictions are decision-support signals, not prescriptions, and the teacher always makes the final judgment in context.",
    )

    doc.add_heading("2. Data", level=1)
    add_labeled_paragraph(
        doc,
        "Dataset route. ",
        "Use the 10,004-row English release of UniversalCEFR on Hugging Face [1] for a standardized data table, then map every row to the official CEFR-SP files [2] to recover the original split and the two annotator labels described by Arase et al. (2022) [3]. Each sample is one independent sentence, not a complete reading text.",
    )

    add_table(
        doc,
        ["Dataset profile", "Audited value"],
        [
            ("Unit of analysis", "One English sentence"),
            ("Number of rows", "10,004"),
            ("Sentence length", "5-34 words; median 14; mean approximately 14.8"),
            ("Fields", "title, lang, source_name, format, category, cefr_level, license, text"),
            ("Public source collections", "SCoRE and Wiki-Auto"),
        ],
        widths=[4.8, 11.0],
    )

    add_table(
        doc,
        ["CEFR level", "A1", "A2", "B1", "B2", "C1", "C2", "Total"],
        [("Sentences", "124", "1,271", "3,305", "3,330", "1,744", "230", "10,004")],
        widths=[2.0, 1.85, 2.0, 2.0, 2.0, 2.0, 1.85, 2.1],
        font_size=10.2,
    )

    add_labeled_paragraph(
        doc,
        "Verified example. ",
        'The dataset contains the sentence "Is that your bike?" with the target label A1. This example will be used in the task definition rather than a fabricated sentence.',
    )
    add_labeled_paragraph(
        doc,
        "Annotation. ",
        "The Hugging Face table contains one CEFR label per sentence. The original CEFR-SP files contain labels from two English-education experts; across all 10,004 matched rows, the Hugging Face label equals the higher of the two levels. In the public SCoRE/Wiki-Auto subset, exact agreement is approximately 41.4%, and every disagreement is within one adjacent CEFR level. The two original labels will be retained for an ambiguity audit rather than treated as identical judgments.",
    )

    add_table(
        doc,
        ["Release/source", "License information to retain", "Project handling"],
        [
            ("UniversalCEFR on Hugging Face", "Records are marked CC BY-NC-SA 4.0", "Retain record-level metadata"),
            ("CEFR-SP: SCoRE", "CC BY-NC-SA 4.0", "Non-commercial educational research"),
            ("CEFR-SP: Wiki-Auto", "CC BY-SA 3.0", "Retain source attribution and share-alike terms"),
        ],
        widths=[4.4, 5.0, 6.4],
        font_size=10.2,
        keep_whole=True,
    )
    add_labeled_paragraph(
        doc,
        "Use restriction. ",
        "The project remains a non-commercial educational research exercise. Source and license metadata will be preserved throughout preprocessing and reporting.",
    )

    doc.add_heading("3. ML Task Alignment", level=1)
    add_table(
        doc,
        ["Interface element", "Definition"],
        [
            ("Input", "One independent English sentence"),
            ("Target", "Exactly one of six labels: A1, A2, B1, B2, C1, C2"),
            ("Inference output", "Six-class probability distribution plus the highest-probability CEFR label"),
            ("Verified training example", 'Text: "Is that your bike?"  |  Target: A1'),
        ],
        widths=[4.1, 11.7],
    )
    add_labeled_paragraph(
        doc,
        "Learning formulation. ",
        "The system performs six-class sentence classification. C1 and C2 remain separate so that the output matches the dataset and the teacher can see the full CEFR scale.",
    )
    add_labeled_paragraph(
        doc,
        "Ordinal structure. ",
        "CEFR labels are ordered: predicting A2 for an A1 sentence is less serious than predicting C2. Training remains a six-class classification task, while evaluation includes quadratic weighted kappa so that larger level gaps receive larger penalties.",
    )

    doc.add_heading("4. Model Plan", level=1)
    add_labeled_paragraph(
        doc,
        "Baseline - majority-class predictor. ",
        "The baseline always predicts the most frequent class in the official training set, B2. It is deliberately simpler than both model families and establishes the minimum result that a learned model must exceed.",
    )
    add_labeled_paragraph(
        doc,
        "Model 1 - word unigram-bigram TF-IDF with multinomial Logistic Regression. ",
        "The vectorizer will represent individual words and short phrases, and the classifier will use balanced class weights. The hypothesis is that lexical choice and short phrase patterns provide strong CEFR signals. To test whether the model is learning topic or source instead of difficulty, the analysis will inspect the largest positive coefficients for each class and report performance separately for SCoRE and Wiki-Auto.",
    )
    add_labeled_paragraph(
        doc,
        "Model 2 - BiLSTM. ",
        "The neural model will use 100-dimensional GloVe word embeddings, a bidirectional LSTM encoder, dropout, class-weighted cross-entropy, and early stopping selected on validation data. The hypothesis is that word order and sequential structure add signals that a bag-of-ngrams representation can miss.",
    )
    add_labeled_paragraph(
        doc,
        "Expected trade-off. ",
        "The sentences are short and the dataset contains only about 10,000 examples, so the BiLSTM may show only a small advantage or may overfit. The comparison is therefore empirical: TF-IDF plus Logistic Regression tests a strong lexical model, while BiLSTM tests whether sequence information adds value under this data budget.",
    )

    doc.add_heading("5. Evaluation Plan", level=1)
    add_labeled_paragraph(
        doc,
        "Official split. ",
        "Do not create a new random 80/10/10 split. Map the Hugging Face rows to the official CEFR-SP split: 7,178 training sentences, 1,366 validation sentences, and 1,460 test sentences. The original split was designed so that highly similar sentences do not appear across training and validation/test. Class proportions are intentionally not identical because the dataset authors reserved additional rare-class examples for validation and test.",
    )
    add_labeled_paragraph(
        doc,
        "Duplicate handling. ",
        "One exact duplicate occurs within the official training split. Remove the repeated copy before fitting, leaving 7,177 unique training sentences; do not move records between splits.",
    )

    add_table(
        doc,
        ["Split", "Mapped rows", "Role"],
        [
            ("Train", "7,178 (7,177 after one duplicate is removed)", "Fit preprocessing and model parameters"),
            ("Validation", "1,366", "Select hyperparameters and early stopping"),
            ("Test", "1,460", "Evaluate once after the model is fixed"),
        ],
        widths=[2.6, 5.0, 8.2],
        keep_whole=True,
    )

    add_labeled_paragraph(doc, "Test-set support. ", "Per-class metrics will always be interpreted together with the following number of test examples:")
    add_table(
        doc,
        ["Class", "A1", "A2", "B1", "B2", "C1", "C2", "Total"],
        [("Sentences", "37", "158", "547", "459", "187", "72", "1,460")],
        widths=[2.0, 1.85, 2.0, 2.0, 2.0, 2.0, 1.85, 2.1],
        font_size=10.2,
    )

    add_labeled_paragraph(
        doc,
        "Leakage control. ",
        "Fit the tokenizer, vocabulary, TF-IDF statistics, label mappings, and any learned preprocessing only on training data. Use validation data for hyperparameter selection and early stopping. Keep the test set untouched until all model choices are fixed.",
    )
    add_labeled_paragraph(
        doc,
        "Primary metric - macro-F1. ",
        "Macro-F1 gives every CEFR class equal weight, so strong performance on the frequent B1/B2 classes cannot hide weak performance on A1 or C2.",
    )
    add_labeled_paragraph(
        doc,
        "Ordinal metric - quadratic weighted kappa. ",
        "This metric uses the order of CEFR levels and penalizes predictions more strongly as the distance from the true level grows. It therefore distinguishes a one-step error such as A1 to A2 from a five-step error such as A1 to C2.",
    )
    add_labeled_paragraph(
        doc,
        "Additional reporting. ",
        "Report accuracy, the confusion matrix, and precision, recall, F1, and support for each class. Estimate 95% confidence intervals with 1,000 stratified bootstrap samples. Interpret A1 F1 especially cautiously because only 37 A1 sentences are present in the test set.",
    )

    doc.add_heading("6. Risks and Responsible Use", level=1)
    add_table(
        doc,
        ["Risk", "Response in this project"],
        [
            ("Label ambiguity", "Two experts do not always agree. Report the approximately 41.4% exact agreement in the public subset, preserve both original labels for audit, and avoid presenting CEFR as perfectly objective."),
            ("Class imbalance", "A1 and C2 are rare. Use class weighting, macro-F1, and per-class support; do not rely on accuracy alone."),
            ("Domain shift", "SCoRE and Wiki-Auto do not fully represent HANU materials or HANU learners. Frame generalization beyond these sources as unverified."),
            ("Topic confounding", "TF-IDF may learn topics or source cues rather than difficulty. Inspect top coefficients and compare performance by source."),
            ("Sentence-to-text limitation", "The model scores sentences independently. Do not infer a single CEFR level for an entire reading text from these predictions."),
            ("Learner variability", "Difficulty also depends on background knowledge, goals, and individual learners. Treat the model output as one signal, not a universal judgment."),
            ("Human control", "Show probabilities and flag uncertain cases. The teacher reviews each suggestion in context and makes the final decision."),
            ("HANU audit", "Before the final project, conduct a qualitative review of about 30 representative sentences from HANU reading materials. Keep these sentences out of training and use them only to examine practical fit and failure modes."),
        ],
        widths=[3.8, 12.0],
        font_size=10.2,
    )

    doc.add_heading("References", level=1)
    references = [
        "[1] UniversalCEFR. cefr_sp_en dataset. https://huggingface.co/datasets/UniversalCEFR/cefr_sp_en",
        "[2] CEFR-SP repository and public split files. https://github.com/yukiar/CEFR-SP/tree/main/CEFR-SP",
        "[3] Arase, Y., Uchida, S., and Kajiwara, T. (2022). CEFR-based Sentence Profile: An Open Dataset for Multilingual Sentence-Level Complexity Assessment. EMNLP 2022. https://aclanthology.org/2022.emnlp-main.416/",
    ]
    for item in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)

    # Document metadata and accessibility-friendly title.
    doc.core_properties.title = "CEFR Based Sentence Difficulty Classification for Teacher Support in Reading Material Review"
    doc.core_properties.subject = "Revised six-part MLA 2026 midterm project outline"
    doc.core_properties.author = "MLA 2026 project team"
    doc.core_properties.keywords = "CEFR, sentence classification, teacher support, TF-IDF, Logistic Regression, BiLSTM"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
