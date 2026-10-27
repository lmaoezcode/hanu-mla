from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "docs" / "reference" / "MLA-2026.docx"
OUTPUT = ROOT / "output" / "docx" / "MLA-2026-revised-outline-simple.docx"


def remove_child_if_present(parent, tag_name):
    child = parent.find(qn(tag_name))
    if child is not None:
        parent.remove(child)


def configure_style(style, size, bold=False, alignment=None, before=0, after=0):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False

    p_format = style.paragraph_format
    p_format.space_before = Pt(before)
    p_format.space_after = Pt(after)
    p_format.line_spacing = 1.0
    if alignment is not None:
        p_format.alignment = alignment

    p_pr = style.element.get_or_add_pPr()
    remove_child_if_present(p_pr, "w:shd")
    remove_child_if_present(p_pr, "w:pBdr")


def add_outline_line(doc, label, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.widow_control = True

    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.name = "Calibri"
    label_run.font.size = Pt(12)
    label_run.font.color.rgb = RGBColor(0, 0, 0)

    text_run = paragraph.add_run(text)
    text_run.bold = False
    text_run.font.name = "Calibri"
    text_run.font.size = Pt(12)
    text_run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_heading(doc, text):
    paragraph = doc.add_paragraph(text, style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_reference(doc, number, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(f"[{number}] {text}")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def build_document():
    doc = Document(REFERENCE)

    body = doc._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)

    configure_style(doc.styles["Normal"], 12, after=7)
    configure_style(
        doc.styles["Title"],
        20,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=7,
    )
    configure_style(
        doc.styles["Subtitle"],
        20,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=13,
    )
    configure_style(doc.styles["Heading 1"], 20, bold=True, before=9, after=7)

    title = doc.add_paragraph(
        "CEFR Based Sentence Difficulty Classification for Teacher Support in Reading Material Review",
        style="Title",
    )
    title.paragraph_format.keep_with_next = True

    subtitle = doc.add_paragraph(
        "Phân loại độ khó câu theo CEFR để hỗ trợ giáo viên rà soát tài liệu đọc",
        style="Subtitle",
    )
    subtitle.paragraph_format.keep_with_next = False

    add_heading(doc, "1. Problem and User")
    add_outline_line(
        doc,
        "Problem",
        "Teachers need to review individual sentences in reading materials and identify language that may be above the target CEFR level.",
    )
    add_outline_line(
        doc,
        "Primary user",
        "English-language teachers at HANU who select or adapt reading materials for a known learner group.",
    )
    add_outline_line(
        doc,
        "Intended use",
        "A teacher provides a reading text. The tool separates it into sentences and predicts a CEFR level for each sentence.",
    )
    add_outline_line(
        doc,
        "Decision supported",
        "The teacher can keep, simplify, replace, or examine a sentence more closely.",
    )
    add_outline_line(
        doc,
        "Human role",
        "The tool shows a predicted level and six class probabilities. The teacher interprets the result in context and makes the final decision.",
    )

    add_heading(doc, "2. Data")
    add_outline_line(
        doc,
        "Dataset",
        "UniversalCEFR/cefr_sp_en on Hugging Face, matched to the official CEFR-SP files [1][2].",
    )
    add_outline_line(
        doc,
        "Unit and size",
        "10,004 independent English sentences. Sentence length is 5-34 words, with a median of 14 and a mean of about 14.8.",
    )
    add_outline_line(
        doc,
        "Fields",
        "title, lang, source_name, format, category, cefr_level, license, and text.",
    )
    add_outline_line(
        doc,
        "Label distribution",
        "A1 124; A2 1,271; B1 3,305; B2 3,330; C1 1,744; C2 230.",
    )
    add_outline_line(
        doc,
        "Public sources",
        "The sentences come from SCoRE and Wiki-Auto [2][3].",
    )
    add_outline_line(doc, "Example", '"Is that your bike?" -> A1.')
    add_outline_line(
        doc,
        "Annotation",
        "CEFR-SP provides two labels from English-education experts. In the matched data, the Hugging Face label is the higher of the two labels.",
    )
    add_outline_line(
        doc,
        "Annotator agreement",
        "Exact agreement in the public subset is about 41.4%, and every difference is within one adjacent CEFR level.",
    )
    add_outline_line(
        doc,
        "Licenses",
        "Hugging Face lists CC BY-NC-SA 4.0. The original repository lists SCoRE as CC BY-NC-SA 4.0 and Wiki-Auto as CC BY-SA 3.0.",
    )
    add_outline_line(
        doc,
        "Planned use",
        "The data will be used for non-commercial educational research, with source and license metadata retained.",
    )

    add_heading(doc, "3. ML Task Alignment")
    add_outline_line(doc, "Task", "Six-class sentence classification.")
    add_outline_line(doc, "Input", "One independent English sentence.")
    add_outline_line(
        doc,
        "Output",
        "Probabilities for A1, A2, B1, B2, C1, and C2, followed by the predicted CEFR level.",
    )
    add_outline_line(doc, "Classes", "A1, A2, B1, B2, C1, and C2.")
    add_outline_line(
        doc,
        "Training example",
        'Input: "Is that your bike?" Target label: A1.',
    )
    add_outline_line(
        doc,
        "Ordinal structure",
        "CEFR levels are ordered. An A1-to-A2 error is smaller than an A1-to-C2 error, so evaluation will also measure error distance.",
    )

    add_heading(doc, "4. Model Plan")
    add_outline_line(
        doc,
        "Baseline",
        "A majority-class predictor that always outputs B2, the most frequent class in the training set.",
    )
    add_outline_line(
        doc,
        "Model 1",
        "Word unigram-bigram TF-IDF with multinomial Logistic Regression and balanced class weights.",
    )
    add_outline_line(
        doc,
        "Model 1 hypothesis",
        "Lexical choices and short phrases provide strong signals of sentence-level CEFR difficulty.",
    )
    add_outline_line(
        doc,
        "Model 1 diagnostic",
        "Inspect top coefficients and compare results for SCoRE and Wiki-Auto to check whether the model is learning topic or source cues.",
    )
    add_outline_line(
        doc,
        "Model 2",
        "A BiLSTM with 100-dimensional GloVe embeddings, dropout, class-weighted cross-entropy, and early stopping.",
    )
    add_outline_line(
        doc,
        "Model 2 hypothesis",
        "Word order and sequential structure may add information that TF-IDF cannot represent.",
    )
    add_outline_line(
        doc,
        "Model 2 limitation",
        "The sentences are short and the dataset has about 10,000 examples, so the BiLSTM may provide only a small gain or may overfit.",
    )

    add_heading(doc, "5. Evaluation Plan")
    add_outline_line(
        doc,
        "Data split",
        "Use the official CEFR-SP split: 7,178 training sentences, 1,366 validation sentences, and 1,460 test sentences.",
    )
    add_outline_line(
        doc,
        "Duplicate handling",
        "Remove one exact duplicate inside the training set before fitting, leaving 7,177 unique training sentences.",
    )
    add_outline_line(
        doc,
        "Split design",
        "The CEFR-SP split keeps highly similar sentences out of both training and validation/test, and includes additional rare-class examples in validation and test.",
    )
    add_outline_line(
        doc,
        "Preprocessing",
        "Fit the tokenizer, vocabulary, and TF-IDF representation using the training set only.",
    )
    add_outline_line(
        doc,
        "Model selection",
        "Use the validation set for hyperparameters and early stopping. Evaluate the test set once after the model choices are fixed.",
    )
    add_outline_line(
        doc,
        "Primary metric",
        "Macro-F1, because the six CEFR classes are imbalanced and each class should contribute equally.",
    )
    add_outline_line(
        doc,
        "Ordinal metric",
        "Quadratic weighted kappa, because CEFR labels are ordered and distant errors are more serious.",
    )
    add_outline_line(
        doc,
        "Additional reporting",
        "Accuracy, confusion matrix, and per-class precision, recall, F1, and support.",
    )
    add_outline_line(
        doc,
        "Uncertainty",
        "Estimate 95% confidence intervals with 1,000 stratified bootstrap samples. Interpret A1 results cautiously because the test set has 37 A1 sentences.",
    )
    add_outline_line(
        doc,
        "Test support",
        "A1 37; A2 158; B1 547; B2 459; C1 187; C2 72.",
    )

    add_heading(doc, "6. Risks and Responsible Use")
    add_outline_line(
        doc,
        "Label ambiguity",
        "Two experts may assign adjacent CEFR levels to the same sentence, so the target label is an informed judgment rather than an absolute fact.",
    )
    add_outline_line(
        doc,
        "Class imbalance",
        "A1 and C2 are rare. The models use class weighting, and evaluation uses macro-F1 with class support.",
    )
    add_outline_line(
        doc,
        "Domain shift",
        "SCoRE and Wiki-Auto may differ from reading materials and learner needs at HANU.",
    )
    add_outline_line(
        doc,
        "Topic confounding",
        "TF-IDF may learn subject matter or source patterns. Coefficients and performance by source will be checked.",
    )
    add_outline_line(
        doc,
        "Sentence-level scope",
        "The model evaluates individual sentences. It does not assign a CEFR level to a complete reading passage.",
    )
    add_outline_line(
        doc,
        "Learner variability",
        "Difficulty may change with a learner's vocabulary, background knowledge, and familiarity with the topic.",
    )
    add_outline_line(
        doc,
        "Human control",
        "Predictions are review suggestions with probabilities. The teacher checks the sentence in context and makes the decision.",
    )
    add_outline_line(
        doc,
        "HANU check",
        "Before the final project, conduct a qualitative review of about 30 representative sentences from HANU reading materials. These sentences are reserved for local review only.",
    )

    add_heading(doc, "References")
    add_reference(
        doc,
        1,
        "UniversalCEFR, cefr_sp_en: https://huggingface.co/datasets/UniversalCEFR/cefr_sp_en",
    )
    add_reference(
        doc,
        2,
        "CEFR-SP repository: https://github.com/yukiar/CEFR-SP/tree/main/CEFR-SP",
    )
    add_reference(
        doc,
        3,
        "Arase et al. (2022), CEFR-SP paper: https://aclanthology.org/2022.emnlp-main.416/",
    )

    doc.core_properties.title = (
        "CEFR Based Sentence Difficulty Classification for Teacher Support in Reading Material Review"
    )
    doc.core_properties.subject = "MLA 2026 project outline"
    doc.core_properties.keywords = "CEFR, sentence classification, teaching support"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
